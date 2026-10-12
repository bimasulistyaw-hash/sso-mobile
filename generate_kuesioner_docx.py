"""
Generate kuesioner_pra_kickoff_kominfo_jombang.docx from the Markdown source.

- Kotak isian (<div height:..px>) -> tabel 1 sel dengan border, tinggi minimal sesuai ukuran S/M/L/XL
  sehingga client bisa langsung mengetik di dalam kotak (kotak ikut membesar jika teks panjang).
- ☐ -> checkbox Word yang bisa diklik (content control).
- "Lainnya: <span>" -> isian teks inline (content control) yang bisa diketik.
- Contoh pengisian (blockquote) -> teks miring abu-abu dengan latar abu muda.

Jalankan ulang skrip ini setiap kali file .md diubah:
    python generate_kuesioner_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).parent
SRC = BASE / "kuesioner_pra_kickoff_kominfo_jombang.md"
OUT = BASE / "kuesioner_pra_kickoff_kominfo_jombang.docx"

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x59, 0x59, 0x59)
INDENT_STEP = 0.6  # cm per level
TEXT_WIDTH_CM = 17.0  # A4 (21cm) - margin kiri/kanan 2cm

# tinggi kotak px (markdown) -> cm (word)
BOX_HEIGHT = {34: 0.9, 70: 1.9, 130: 3.4, 180: 4.8}

W14 = 'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'

INLINE_RE = re.compile(
    r"(\*\*\*(?P<bi>.+?)\*\*\*"
    r"|\*\*(?P<b>.+?)\*\*"
    r"|(?<!\*)\*(?!\*)(?P<i>.+?)(?<!\*)\*(?!\*)"
    r"|`(?P<c>[^`]+)`"
    r"|(?P<span><span[^>]*></span>))"
)
DIV_RE = re.compile(r'<div style="[^"]*width:\s*(\d+)%;\s*height:\s*(\d+)px[^"]*"></div>')
EMOJI_RE = re.compile(r"[\U0001F300-\U0001FAFF]\s*")


# ---------------------------------------------------------------- helpers
def set_cell_border(table, color="7F7F7F", size=6):
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        + "".join(
            f'<w:{e} w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
            for e in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders>"
    )
    tbl_pr.append(borders)


def set_table_indent(table, cm):
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{int(cm * 567)}" w:type="dxa"/>'))


def shade_paragraph(p, fill):
    p._p.get_or_add_pPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def bottom_border(p, color="BFBFBF"):
    p._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{color}"/></w:pBdr>'
        )
    )


def add_checkbox(p):
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")} {W14}>'
        '<w:sdtPr>'
        '<w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic"/></w:rPr>'
        '<w14:checkbox>'
        '<w14:checked w14:val="0"/>'
        '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/>'
        '</w14:checkbox>'
        '</w:sdtPr>'
        '<w:sdtContent><w:r>'
        '<w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic"/></w:rPr>'
        '<w:t>\u2610</w:t>'
        '</w:r></w:sdtContent>'
        '</w:sdt>'
    )
    p._p.append(sdt)


def add_inline_field(p):
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}>'
        '<w:sdtPr><w:showingPlcHdr/><w:text/></w:sdtPr>'
        '<w:sdtContent><w:r><w:rPr><w:color w:val="808080"/></w:rPr>'
        '<w:t xml:space="preserve">........................................................</w:t>'
        '</w:r></w:sdtContent>'
        '</w:sdt>'
    )
    p._p.append(sdt)


def add_text(p, text, bold=False, italic=False, code=False, color=None, size=None):
    """Tambah teks; setiap ☐ diganti checkbox yang bisa diklik."""
    parts = text.split("\u2610")
    for idx, part in enumerate(parts):
        if idx > 0:
            add_checkbox(p)
        if part:
            r = p.add_run(part)
            r.bold = bold or None
            r.italic = italic or None
            if code:
                r.font.name = "Consolas"
            if color:
                r.font.color.rgb = color
            if size:
                r.font.size = Pt(size)


def add_inline(p, text, bold=False, italic=False, color=None, size=None):
    text = text.replace("&nbsp;", "\u00a0").replace("&amp;", "&")
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            add_text(p, text[pos:m.start()], bold, italic, color=color, size=size)
        if m.group("bi") is not None:
            add_inline(p, m.group("bi"), True, True, color, size)
        elif m.group("b") is not None:
            add_inline(p, m.group("b"), True, italic, color, size)
        elif m.group("i") is not None:
            add_inline(p, m.group("i"), bold, True, color, size)
        elif m.group("c") is not None:
            add_text(p, m.group("c"), bold, italic, code=True, color=color, size=size)
        elif m.group("span") is not None:
            add_inline_field(p)
        pos = m.end()
    if pos < len(text):
        add_text(p, text[pos:], bold, italic, color=color, size=size)


def para(doc, indent_cm=0.0, hanging_cm=0.0, before=0, after=3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(indent_cm)
    if hanging_cm:
        pf.first_line_indent = Cm(-hanging_cm)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return p


def add_box(doc, indent_cm, width_pct, height_px):
    avail = TEXT_WIDTH_CM - indent_cm
    width = Cm(avail * width_pct / 100)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_cell_border(table)
    set_table_indent(table, indent_cm)
    row = table.rows[0]
    row.height = Cm(BOX_HEIGHT.get(height_px, height_px * 0.026))
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.width = width
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    # spasi kecil setelah kotak
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(6)


# ---------------------------------------------------------------- main
def build():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # footer nomor halaman
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Kuesioner Pra-Kick-Off JOSS Mobile — Halaman ")
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY
    for xml in (
        f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>',
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>',
        f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>',
    ):
        run = fp.add_run()
        run.font.size = Pt(9)
        run._r.append(parse_xml(xml))

    header_done = False
    for raw in lines:
        line = EMOJI_RE.sub("", raw.rstrip())
        stripped = line.strip()
        spaces = len(line) - len(line.lstrip(" "))

        if not stripped:
            continue

        # --- judul dokumen
        if stripped.startswith("# "):
            p = para(doc, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped[2:], bold=True, color=NAVY, size=18)
            continue
        if stripped.startswith("## ") and not header_done:
            header_done = True
            p = para(doc, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped[3:], bold=True, color=NAVY, size=14)
            continue
        if stripped.startswith("**Dinas") and len(doc.paragraphs) < 4:
            p = para(doc, after=6)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped, size=12)
            bottom_border(p, "1F3A5F")
            continue

        # --- garis pemisah
        if stripped == "---":
            p = para(doc, after=6)
            p.paragraph_format.line_spacing = Pt(4)
            bottom_border(p)
            continue

        # --- header bagian (pita biru)
        if stripped.startswith("## "):
            p = para(doc, before=12, after=6)
            shade_paragraph(p, "1F3A5F")
            add_inline(p, " " + stripped[3:], bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=12)
            p.paragraph_format.keep_with_next = True
            continue
        if stripped.startswith("### "):
            p = para(doc, before=6, after=4)
            add_inline(p, stripped[4:], bold=True, color=NAVY, size=13)
            p.paragraph_format.keep_with_next = True
            continue
        if stripped.startswith("#### "):
            p = para(doc, before=8, after=4)
            add_inline(p, stripped[5:], bold=True, color=NAVY, size=12)
            p.paragraph_format.keep_with_next = True
            continue

        # --- kotak isian
        m = DIV_RE.search(stripped)
        if m and stripped.startswith("<div"):
            add_box(doc, (spaces // 2) * INDENT_STEP, int(m.group(1)), int(m.group(2)))
            continue

        # --- contoh pengisian (blockquote)
        if stripped.startswith(">"):
            content = stripped[1:]
            content = content[1:] if content.startswith(" ") else content
            inner_sp = len(content) - len(content.lstrip(" "))
            content = content.strip()
            lvl = inner_sp // 2
            base = 0.6 + lvl * INDENT_STEP
            if content.startswith("* "):
                p = para(doc, indent_cm=base + 0.4, hanging_cm=0.4, after=0)
                add_text(p, "•\t", color=GRAY, size=10)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(base + 0.4))
                content = content[2:]
            else:
                p = para(doc, indent_cm=base, after=0)
            shade_paragraph(p, "F2F2F2")
            add_inline(p, content, italic=True, color=GRAY, size=10)
            continue

        if stripped == "* *Contoh Pengisian:*":
            p = para(doc, indent_cm=0, before=2, after=1)
            add_inline(p, "Contoh Pengisian:", bold=True, italic=True, color=GRAY, size=10)
            p.paragraph_format.keep_with_next = True
            continue

        # --- bullet
        if stripped.startswith("* "):
            lvl = spaces // 2
            left = (lvl + 1) * INDENT_STEP
            p = para(doc, indent_cm=left, hanging_cm=0.4, after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(left))
            add_text(p, ("•" if lvl == 0 else "◦") + "\t")
            add_inline(p, stripped[2:])
            p.paragraph_format.keep_with_next = True
            continue

        # --- daftar bernomor
        om = re.match(r"(\d+)\.\s+(.*)", stripped)
        if om:
            left = (spaces // 2 + 1) * INDENT_STEP
            p = para(doc, indent_cm=left, hanging_cm=0.5, after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(left))
            add_text(p, om.group(1) + ".\t")
            add_inline(p, om.group(2))
            continue

        # --- paragraf biasa / baris lanjutan (checkbox, label)
        p = para(doc, indent_cm=(spaces // 2) * INDENT_STEP, after=2)
        add_inline(p, stripped)

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
